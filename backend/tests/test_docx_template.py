from __future__ import annotations

import asyncio

from openpyxl import Workbook
from reportlab.pdfgen import canvas

from app.services.export.word import _write_docx
from app.services.template.docx import DocxTemplateService


def test_docx_template_service_extracts_placeholders(tmp_path) -> None:
    template_path = tmp_path / "daily_report_template.docx"
    _write_docx(
        template_path,
        """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:r><w:t>{{title}}</w:t></w:r></w:p>
    <w:p><w:r><w:t>日期：{{date}}</w:t></w:r></w:p>
    <w:p><w:r><w:t>{{summary}}</w:t></w:r></w:p>
    <w:p><w:r><w:t>{{completed_tasks}}</w:t></w:r></w:p>
    <w:p><w:r><w:t>{{custom_fields.owner}}</w:t></w:r></w:p>
  </w:body>
</w:document>
""",
    )

    async def run() -> None:
        service = DocxTemplateService()
        result = await service.parse_template(str(template_path))

        assert result.template_type == "daily"
        assert result.fields == [
            "title",
            "date",
            "summary",
            "completed_tasks",
            "custom_fields.owner",
        ]
        assert result.raw_content["placeholder_count"] == 5

    asyncio.run(run())


def test_xlsx_template_service_extracts_placeholders(tmp_path) -> None:
    template_path = tmp_path / "daily_report_template.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = "{{title}}"
    worksheet["A2"] = "{{summary}}"
    worksheet["A3"] = "{{任务完成情况}}"
    worksheet["A4"] = "{{owner}}"
    workbook.save(template_path)

    async def run() -> None:
        service = DocxTemplateService()
        result = await service.parse_template(str(template_path))

        assert result.template_type == "daily"
        assert result.fields == ["title", "summary", "completed_tasks", "owner"]
        assert result.raw_content["placeholder_count"] == 4
        assert result.raw_content["source"] == "xlsx"

    asyncio.run(run())


def test_pdf_template_service_extracts_placeholders(tmp_path) -> None:
    template_path = tmp_path / "daily_report_template.pdf"
    _create_template_pdf(template_path)

    async def run() -> None:
        service = DocxTemplateService()
        result = await service.parse_template(str(template_path))

        assert result.template_type == "daily"
        assert result.fields == ["title", "completed_tasks", "owner"]
        assert result.raw_content["placeholder_count"] == 3
        assert result.raw_content["source"] == "pdf"

    asyncio.run(run())


def _create_template_pdf(template_path) -> None:
    pdf = canvas.Canvas(str(template_path), pageCompression=0)
    pdf.drawString(72, 720, "{{title}}")
    pdf.drawString(72, 700, "{{completed_tasks}}")
    pdf.drawString(72, 680, "{{owner}}")
    pdf.save()
